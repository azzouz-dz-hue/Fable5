"""Analyse des relevés téléchargés et génération des exports."""

from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from bankextract.export import export_csv, export_excel, export_ofx, transactions_dataframe
from bankextract.models import Account, Transaction
from bankextract.parsers import parse_statement
from bankextract.parsers.csv_statement import map_columns, parse_csv_statement

COMPTE = "bna:0010098765432109"


def test_correspondance_des_colonnes_privilegie_la_correspondance_exacte():
    """« Date » ne doit pas absorber « Date valeur »."""
    mapping = map_columns(["Date Opération", "Date valeur", "Libellé", "Débit", "Crédit"])
    assert mapping["date"] == "Date Opération"
    assert mapping["value_date"] == "Date valeur"


def test_correspondance_des_colonnes_montant_et_sens():
    mapping = map_columns(["DATE", "LIBELLE", "MONTANT", "SENS"])
    assert mapping == {"date": "DATE", "label": "LIBELLE", "amount": "MONTANT", "sense": "SENS"}


def test_lecture_csv_colonnes_debit_credit(tmp_path):
    chemin = tmp_path / "releve.csv"
    chemin.write_text(
        "Date opération;Date valeur;Libellé;Débit;Crédit;Solde\n"
        "04/03/2026;05/03/2026;VIR RECU CLIENT;;1 250 000,00;4 235 890,45\n"
        "03/03/2026;03/03/2026;REGLEMENT FOURNISSEUR;845 300,50;;2 985 890,45\n"
        "TOTAL;;;845 300,50;1 250 000,00;\n",
        encoding="utf-8",
    )

    transactions = parse_csv_statement(chemin, COMPTE)

    assert len(transactions) == 2, "la ligne de total doit être ignorée"
    assert transactions[0].amount == Decimal("1250000.00")
    assert transactions[0].value_date == date(2026, 3, 5)
    assert transactions[1].amount == Decimal("-845300.50")


def test_lecture_csv_colonne_sens(tmp_path):
    chemin = tmp_path / "releve.csv"
    chemin.write_text(
        "DATE,LIBELLE,MONTANT,SENS\n04/03/2026,FRAIS,1200.00,D\n05/03/2026,VIR,900.00,C\n",
        encoding="utf-8",
    )

    transactions = parse_csv_statement(chemin, COMPTE)

    assert [t.amount for t in transactions] == [Decimal("-1200.00"), Decimal("900.00")]


def test_lecture_csv_encodage_windows(tmp_path):
    """Les portails exportent encore couramment en cp1252."""
    chemin = tmp_path / "releve.csv"
    chemin.write_bytes(
        "Date;Libellé;Débit;Crédit\n04/03/2026;RÈGLEMENT DÉPÔT;1 200,00;\n".encode("cp1252")
    )

    transactions = parse_csv_statement(chemin, COMPTE)

    assert transactions[0].label == "RÈGLEMENT DÉPÔT"


def test_lecture_csv_sans_colonne_date(tmp_path):
    chemin = tmp_path / "releve.csv"
    chemin.write_text("A;B\n1;2\n", encoding="utf-8")

    with pytest.raises(ValueError, match="date"):
        parse_csv_statement(chemin, COMPTE)


def test_format_non_pris_en_charge(tmp_path):
    chemin = tmp_path / "releve.docx"
    chemin.write_text("x")

    with pytest.raises(ValueError, match="non pris en charge"):
        parse_statement(chemin, COMPTE)


# ------------------------------------------------------------------ exports


@pytest.fixture
def ecritures():
    return [
        Transaction(
            account_key=COMPTE,
            date=date(2026, 3, 4),
            value_date=date(2026, 3, 5),
            label="VIR RECU CLIENT",
            amount=Decimal("1250000.00"),
            balance_after=Decimal("4235890.45"),
            reference="VIR26030401",
        ),
        Transaction(
            account_key=COMPTE,
            date=date(2026, 3, 3),
            label="REGLEMENT FOURNISSEUR",
            amount=Decimal("-845300.50"),
        ),
        Transaction(
            account_key="bna:0010098765432241",
            date=date(2026, 3, 3),
            label="FRAIS SWIFT",
            amount=Decimal("-120.00"),
            currency="EUR",
        ),
    ]


def test_tableau_normalise_separe_debit_et_credit(ecritures):
    frame = transactions_dataframe(ecritures)

    assert list(frame.columns)[:5] == [
        "Banque", "Compte", "Date opération", "Date valeur", "Libellé"
    ]
    credit = frame[frame["Libellé"] == "VIR RECU CLIENT"].iloc[0]
    assert credit["Crédit"] == Decimal("1250000.00") and credit["Débit"] is None


def test_export_csv_lisible_par_excel_francophone(tmp_path, ecritures):
    chemin = export_csv(ecritures, tmp_path / "sortie.csv")
    contenu = chemin.read_text(encoding="utf-8-sig")

    lignes = contenu.splitlines()
    assert lignes[0].startswith("Banque;Compte;Date opération")
    assert "04/03/2026" in contenu, "les dates doivent être au format français"


def test_export_excel_une_feuille_par_compte(tmp_path, ecritures):
    compte = Account(bank="bna", number="00100 987654321 09", balance=Decimal("4235890.45"))
    chemin = export_excel(
        ecritures,
        tmp_path / "sortie.xlsx",
        account_labels={COMPTE: "00100 987654321 09"},
        accounts=[compte],
    )

    classeur = openpyxl.load_workbook(chemin)
    assert classeur.sheetnames[0] == "Synthèse"
    assert len(classeur.sheetnames) == 3  # synthèse + deux comptes

    entetes = [cellule.value for cellule in classeur["Synthèse"][1]]
    assert "Total débit" in entetes and "Variation nette" in entetes


def test_export_excel_sans_ecriture(tmp_path):
    chemin = export_excel([], tmp_path / "vide.xlsx")

    assert openpyxl.load_workbook(chemin).sheetnames == ["Synthèse", "Écritures"]


def test_export_ofx_contient_chaque_ecriture(tmp_path, ecritures):
    chemin = export_ofx(ecritures, tmp_path / "sortie.ofx", account_number="987654321")
    contenu = chemin.read_text(encoding="utf-8")

    assert contenu.startswith("OFXHEADER:100")
    assert contenu.count("<STMTTRN>") == len(ecritures)
    assert "<TRNTYPE>CREDIT" in contenu and "<TRNTYPE>DEBIT" in contenu
    assert "<ACCTID>987654321" in contenu


def test_nom_de_feuille_tronque_et_assaini(tmp_path):
    """Excel refuse les noms de plus de 31 caractères et certains symboles."""
    longue = [
        Transaction(
            account_key="bna:X",
            date=date(2026, 3, 4),
            label="OP",
            amount=Decimal(1),
        )
    ]
    chemin = export_excel(
        longue, tmp_path / "s.xlsx", account_labels={"bna:X": "A" * 40 + "/[test]"}
    )

    nom = openpyxl.load_workbook(chemin).sheetnames[1]
    assert len(nom) <= 31 and "/" not in nom and "[" not in nom
